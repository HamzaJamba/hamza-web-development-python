import io
from django.shortcuts import render, get_object_or_404
from django.http import HttpResponse
from django.views import View
from django.db.models import Count, Sum

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from accounts.mixins import AdminRequiredMixin, StaffRequiredMixin
from doctors.models import Doctor
from patients.models import Patient
from appointments.models import Appointment
from billing.models import Invoice
from prescriptions.models import Prescription

BRAND_COLOR = colors.HexColor("#2563EB")


def _pdf_header(elements, styles, title, subtitle=""):
    title_style = ParagraphStyle("TitleStyle", parent=styles["Title"], fontSize=18, textColor=BRAND_COLOR)
    elements.append(Paragraph("City Care Hospital", title_style))
    elements.append(Paragraph(title, styles["Heading2"]))
    if subtitle:
        elements.append(Paragraph(subtitle, styles["Normal"]))
    elements.append(Spacer(1, 0.5 * cm))


def _styled_table(data, col_widths=None):
    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), BRAND_COLOR),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return table


class ReportsHomeView(StaffRequiredMixin, View):
    def get(self, request):
        return render(request, "reports/reports_home.html", {"doctors": Doctor.objects.filter(is_active=True)})


class PatientsReportPDFView(StaffRequiredMixin, View):
    def get(self, request):
        patients = Patient.objects.all()
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        styles = getSampleStyleSheet()
        elements = []
        _pdf_header(elements, styles, "Patients Report")

        data = [["Patient ID", "Name", "Age", "Gender", "Phone", "Blood Group", "Status"]]
        for p in patients:
            data.append([p.patient_id, p.full_name, p.age, p.get_gender_display(), p.phone,
                         p.blood_group or "-", "Active" if p.is_active else "Inactive"])
        elements.append(_styled_table(data))
        doc.build(elements)

        response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
        response["Content-Disposition"] = 'attachment; filename="patients_report.pdf"'
        return response


class AppointmentsReportPDFView(StaffRequiredMixin, View):
    def get(self, request):
        appointments = Appointment.objects.select_related("patient", "doctor__user").order_by("-date", "-time")
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        styles = getSampleStyleSheet()
        elements = []
        _pdf_header(elements, styles, "Appointments Report")

        data = [["Date", "Time", "Patient", "Doctor", "Reason", "Status"]]
        for a in appointments:
            data.append([str(a.date), a.time.strftime("%H:%M"), a.patient.full_name,
                         str(a.doctor), a.reason or "-", a.get_status_display()])
        elements.append(_styled_table(data))
        doc.build(elements)

        response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
        response["Content-Disposition"] = 'attachment; filename="appointments_report.pdf"'
        return response


class DoctorWiseReportPDFView(StaffRequiredMixin, View):
    def get(self, request):
        doctors = Doctor.objects.select_related("user").annotate(
            total_appts=Count("appointments"), completed=Count("appointments", filter=None)
        )
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        styles = getSampleStyleSheet()
        elements = []
        _pdf_header(elements, styles, "Doctor-wise Appointments Summary")

        data = [["Doctor", "Specialization", "Total Appointments", "Completed", "Scheduled", "Cancelled"]]
        for d in doctors:
            qs = d.appointments.all()
            data.append([
                str(d), d.specialization, qs.count(),
                qs.filter(status="COMPLETED").count(),
                qs.filter(status="SCHEDULED").count(),
                qs.filter(status="CANCELLED").count(),
            ])
        elements.append(_styled_table(data))
        doc.build(elements)

        response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
        response["Content-Disposition"] = 'attachment; filename="doctor_wise_report.pdf"'
        return response


class BillingReportPDFView(AdminRequiredMixin, View):
    def get(self, request):
        invoices = Invoice.objects.select_related("patient").all()
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        styles = getSampleStyleSheet()
        elements = []
        _pdf_header(elements, styles, "Billing Report")

        data = [["Invoice #", "Patient", "Total", "Paid", "Balance", "Status", "Date"]]
        total_collected = 0
        for inv in invoices:
            data.append([inv.invoice_number, inv.patient.full_name, str(inv.total_amount),
                         str(inv.amount_paid), str(inv.balance), inv.get_status_display(), str(inv.date)])
            total_collected += inv.amount_paid
        elements.append(_styled_table(data))
        elements.append(Spacer(1, 0.5 * cm))
        elements.append(Paragraph(f"Total Collected: Rs. {total_collected}", styles["Heading3"]))
        doc.build(elements)

        response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
        response["Content-Disposition"] = 'attachment; filename="billing_report.pdf"'
        return response


class PatientMedicalReportPDFView(StaffRequiredMixin, View):
    """Full medical report for a single patient: history + prescriptions."""
    def get(self, request, patient_id):
        patient = get_object_or_404(Patient, pk=patient_id)
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        styles = getSampleStyleSheet()
        elements = []
        _pdf_header(elements, styles, "Patient Medical Report", f"{patient.full_name} ({patient.patient_id})")

        elements.append(Paragraph("Medical History", styles["Heading2"]))
        data = [["Date", "Doctor", "Diagnosis", "Notes"]]
        for h in patient.medical_history.select_related("doctor__user"):
            data.append([str(h.date), str(h.doctor) if h.doctor else "-", h.diagnosis, h.notes or "-"])
        elements.append(_styled_table(data))
        elements.append(Spacer(1, 0.6 * cm))

        elements.append(Paragraph("Prescriptions", styles["Heading2"]))
        for pres in patient.prescriptions.select_related("doctor__user").prefetch_related("medicines"):
            elements.append(Paragraph(f"{pres.date} — Dr. {pres.doctor} — {pres.diagnosis or 'General'}", styles["Heading4"]))
            med_data = [["Medicine", "Dosage", "Duration", "Instructions"]]
            for m in pres.medicines.all():
                med_data.append([m.name, m.dosage, m.duration, m.instructions or "-"])
            elements.append(_styled_table(med_data))
            elements.append(Spacer(1, 0.3 * cm))

        doc.build(elements)
        response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="medical_report_{patient.patient_id}.pdf"'
        return response


def _excel_response(filename, header, rows, sheet_title="Report"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title

    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for col, title in enumerate(header, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = max_len + 4

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


class PatientsExcelExportView(StaffRequiredMixin, View):
    def get(self, request):
        patients = Patient.objects.all()
        header = ["Patient ID", "Name", "Age", "Gender", "Phone", "Blood Group", "Address", "Status"]
        rows = [
            [p.patient_id, p.full_name, p.age, p.get_gender_display(), p.phone, p.blood_group or "-",
             p.address, "Active" if p.is_active else "Inactive"]
            for p in patients
        ]
        return _excel_response("patients.xlsx", header, rows, "Patients")


class AppointmentsExcelExportView(StaffRequiredMixin, View):
    def get(self, request):
        appointments = Appointment.objects.select_related("patient", "doctor__user").all()
        header = ["Date", "Time", "Patient", "Doctor", "Reason", "Status"]
        rows = [
            [str(a.date), a.time.strftime("%H:%M"), a.patient.full_name, str(a.doctor), a.reason or "-", a.get_status_display()]
            for a in appointments
        ]
        return _excel_response("appointments.xlsx", header, rows, "Appointments")


class BillingExcelExportView(AdminRequiredMixin, View):
    def get(self, request):
        invoices = Invoice.objects.select_related("patient").all()
        header = ["Invoice #", "Patient", "Consultation Fee", "Other Charges", "Total", "Paid", "Balance", "Status", "Date"]
        rows = [
            [inv.invoice_number, inv.patient.full_name, float(inv.consultation_fee), float(inv.other_charges),
             float(inv.total_amount), float(inv.amount_paid), float(inv.balance), inv.get_status_display(), str(inv.date)]
            for inv in invoices
        ]
        return _excel_response("billing_report.xlsx", header, rows, "Billing")
