import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def to_xlsx(report):
    wb = Workbook()
    ws = wb.active
    ws.title = report["title"][:31]
    ws.append([report["title"]])
    ws["A1"].font = Font(size=14, bold=True)
    ws.append([report["subtitle"]])
    ws.append([])
    ws.append(report["columns"])
    head = 4
    for c in ws[head]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0F2A2E")
        c.alignment = Alignment(horizontal="center")
    for r in report["rows"]:
        ws.append(r)
    if report["summary"]:
        ws.append([])
        for k, v in report["summary"].items():
            ws.append([k, v])
            ws.cell(ws.max_row, 1).font = Font(bold=True)
    for i, col in enumerate(report["columns"], 1):
        width = max([len(str(col))] + [len(str(r[i - 1])) for r in report["rows"] if len(r) >= i]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(width, 45)
    ws.freeze_panes = ws.cell(head + 1, 1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(report):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=28, bottomMargin=28,
                            title=report["title"])
    st = getSampleStyleSheet()
    story = [Paragraph(report["title"], st["Title"]), Paragraph(report["subtitle"], st["Normal"]), Spacer(1, 12)]
    small = st["BodyText"].clone("small", fontSize=7, leading=9)
    data = [[Paragraph(f"<b>{c}</b>", small) for c in report["columns"]]]
    data += [[Paragraph(str(v), small) for v in r] for r in report["rows"]] or [["No data for the selected period."]]
    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#d9e2e4")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#9aa9ad")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f6f7")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(t)
    if report["summary"]:
        story.append(Spacer(1, 12))
        for k, v in report["summary"].items():
            story.append(Paragraph(f"<b>{k}:</b> {v}", st["Normal"]))
    doc.build(story)
    return buf.getvalue()
